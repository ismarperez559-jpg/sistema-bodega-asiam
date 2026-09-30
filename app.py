import os
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Sistema de Bodega - ASIAM", page_icon="⚙", layout="wide"
)

# Búsqueda automática del archivo Excel en la carpeta
archivos = [f for f in os.listdir(".") if f.endswith(".xlsx")]
EXCEL_FILE = archivos[0] if archivos else "inventario.xlsx"


@st.cache_data
def cargar_datos():
  if os.path.exists(EXCEL_FILE):
    try:
      df = pd.read_excel(EXCEL_FILE, sheet_name="BUSCADOR", header=11)
      df = df.loc[:, ~df.columns.str.contains("^Unnamed")]
      df = df.dropna(subset=["Id", "PRODUCTO"], how="all")
      df.columns = df.columns.astype(str).str.strip()
      return df
    except Exception as e:
      try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=0)
        df = df.loc[:, ~df.columns.str.contains("^Unnamed")]
        df.columns = df.columns.astype(str).str.strip()
        return df
      except:
        return pd.DataFrame()
  else:
    return pd.DataFrame()


df_inventario = cargar_datos()

st.title("📦 Sistema de Gestión y Cruce de Repuestos - ASIAM")
st.sidebar.header("Menú de Bodega")
opcion = st.sidebar.radio(
    "Seleccione una opción:",
    [
        "Inventario y Cruce de Repuestos",
        "Registrar Nuevo Repuesto",
        "Impresión Zebra (ZPL)",
    ],
)

if opcion == "Inventario y Cruce de Repuestos":
  st.header("🔍 Consulta y Cruce Avanzado de Repuestos")

  if df_inventario.empty:
    st.error(f"⚠️ No se pudo leer el archivo '{EXCEL_FILE}'.")
  else:
    st.sidebar.markdown("---")
    st.sidebar.subheader("Filtros de Bodega")

    if "Marca" in df_inventario.columns:
      marcas_disponibles = ["Todas"] + sorted(
          df_inventario["Marca"].dropna().astype(str).unique().tolist()
      )
      marca_sel = st.sidebar.selectbox("Filtrar por Marca", marcas_disponibles)
    else:
      marca_sel = "Todas"

    if "Categoria Primaria" in df_inventario.columns:
      categorias_disponibles = ["Todas"] + sorted(
          df_inventario["Categoria Primaria"]
          .dropna()
          .astype(str)
          .unique()
          .tolist()
      )
      cat_sel = st.sidebar.selectbox(
          "Filtrar por Categoría", categorias_disponibles
      )
    else:
      cat_sel = "Todas"

    df_filtrado = df_inventario.copy()
    if marca_sel != "Todas":
      df_filtrado = df_filtrado[
          df_filtrado["Marca"].astype(str) == marca_sel
      ]
    if cat_sel != "Todas":
      df_filtrado = df_filtrado[
          df_filtrado["Categoria Primaria"].astype(str) == cat_sel
      ]

    busqueda = st.text_input(
        "Buscar repuesto (ID, Producto, OEM, Código Asiam o Cruces):"
    )

    if busqueda:
      df_filtrado = df_filtrado[
          df_filtrado.astype(str)
          .apply(lambda row: row.str.contains(busqueda, case=False).any(), axis=1)
      ]

    st.info(f"Mostrando {len(df_filtrado)} repuestos encontrados:")
    st.dataframe(df_filtrado, use_container_width=True)

elif opcion == "Registrar Nuevo Repuesto":
  st.header("➕ Ingreso de Nuevo Repuesto a Bodega")

  with st.form("form_nuevo_repuesto"):
    col1, col2 = st.columns(2)
    with col1:
      prod_nombre = st.text_input("Nombre del Producto / Descripción").upper()
      marca_rep = st.text_input(
          "Marcas Compatibles (Ej: MG ZS / CHERY TIGGO 2)"
      ).upper()
      cat_rep = st.selectbox(
          "Categoría Primaria",
          [
              "SUSPENSION",
              "MOTOR",
              "LUCES",
              "FRENOS",
              "FILTROS",
              "SENSOR",
              "CARROCERIA",
              "VARIOS",
          ],
      )
    with col2:
      oem_rep = st.text_input("Código Producto (OEM)").upper()
      interno_rep = st.text_input("Código Interno (Opcional)").upper()

    submitted = st.form_submit_button("💾 Guardar y Generar Código Asiam")

    if submitted:
      if not prod_nombre or not oem_rep:
        st.warning(
            "⚠️ Por favor ingresa al menos el Nombre del Producto y el Código"
            " OEM."
        )
      else:
        # Generar ID aleatorio de bodega
        import random

        nuevo_id = random.randint(3300000, 3999999)

        # Autogenerar Código Asiam secuencial según categoría
        prefijo_map = {
            "SUSPENSION": "ABS",
            "MOTOR": "ACE",
            "LUCES": "LUC",
            "FRENOS": "FRE",
            "FILTROS": "FIL",
            "SENSOR": "SEN",
            "CARROCERIA": "CAR",
            "VARIOS": "VAR",
        }
        pref = prefijo_map.get(cat_rep, "ASI")

        # Filtrar cuántos hay con ese prefijo para calcular el consecutivo
        if not df_inventario.empty and "Código Asiam" in df_inventario.columns:
          existentes = df_inventario[
              df_inventario["Código Asiam"].astype(str).str.startswith(pref)
          ]
          siguiente_num = len(existentes) + 1
        else:
          siguiente_num = 1

        nuevo_asiam = f"{pref}-{siguiente_num:04d}"
        nuevo_cruces = f"{nuevo_asiam} / {oem_rep}"
        nuevo_barras = f"*{nuevo_asiam}*"

        # Nuevo registro
        nuevo_registro = {
            "Id": nuevo_id,
            "Marca": marca_rep,
            "PRODUCTO": prod_nombre,
            "Categoria Primaria": cat_rep,
            "Código Producto": oem_rep,
            "Códigos Interno": (
                interno_rep if interno_rep else str(random.randint(3820000, 3999999))
            ),
            "Código Asiam": nuevo_asiam,
            "Cruces Disponibles": nuevo_cruces,
            "Código Barras": nuevo_barras,
        }

        # Guardar en el Excel existente respetando la estructura
        try:
          import openpyxl

          # Cargar libro y agregar a la hoja BUSCADOR
          book = openpyxl.load_workbook(EXCEL_FILE)
          sheet = book["BUSCADOR"]

          # Encontrar la última fila con datos
          next_row = sheet.max_row + 1
          sheet.cell(row=next_row, column=2, value=nuevo_id)
          sheet.cell(row=next_row, column=3, value=marca_rep)
          sheet.cell(row=next_row, column=4, value=prod_nombre)
          sheet.cell(row=next_row, column=5, value=cat_rep)
          sheet.cell(row=next_row, column=6, value=oem_rep)
          sheet.cell(
              row=next_row,
              column=7,
              value=int(nuevo_registro["Códigos Interno"]),
          )
          sheet.cell(row=next_row, column=8, value=nuevo_asiam)
          sheet.cell(row=next_row, column=9, value=nuevo_cruces)
          sheet.cell(row=next_row, column=10, value=nuevo_barras)

          book.save(EXCEL_FILE)
          st.success(
              f"🎉 ¡Repuesto registrado con éxito! Código Asiam asignado:"
              f" **{nuevo_asiam}**"
          )
          st.balloons()
        except Exception as ex:
          st.error(f"Error al guardar en el archivo Excel: {ex}")

elif opcion == "Impresión Zebra (ZPL)":
  st.header("🖨️ Búsqueda y Emisión Directa para Impresora Zebra")

  if not df_inventario.empty:
    texto_buscar = st.text_input(
        "🔎 Escribe parte del producto, código Asiam o ID para etiquetar:"
    )

    df_etiquetas = df_inventario.copy()
    if texto_buscar:
      df_etiquetas = df_etiquetas[
          df_etiquetas.astype(str).apply(
              lambda row: row.str.contains(texto_buscar, case=False).any(), axis=1
          )
      ]

    if df_etiquetas.empty:
      st.warning("No se encontraron repuestos con ese criterio de búsqueda.")
    else:
      opciones_select = []
      for idx, row in df_etiquetas.iterrows():
        id_i = str(row.get("Id", ""))
        asiam_i = str(row.get("Código Asiam", ""))
        prod_i = str(row.get("PRODUCTO", ""))
        oem_i = str(row.get("Código Product", ""))
        etiqueta_txt = (
            f"ID: {id_i} | Asiam: {asiam_i} | {prod_i} | OEM: {oem_i}"
        )
        opciones_select.append((id_i, etiqueta_txt))

      seleccion_usuario = st.selectbox(
          "Selecciona el repuesto a etiquetar:",
          options=opciones_select,
          format_func=lambda x: x[1],
      )

      if seleccion_usuario:
        id_elegido = seleccion_usuario[0]
        item = df_inventario[
            df_inventario["Id"].astype(str) == str(id_elegido)
        ].iloc[0]

        id_val = str(item.get("Id", ""))
        marca_val = str(item.get("Marca", ""))
        prod_val = str(item.get("PRODUCTO", ""))
        oem_val = str(item.get("Código Product", ""))
        asiam_val = str(item.get("Código Asiam", ""))

        col1, col2 = st.columns(2)
        with col1:
          cantidad_copias = st.number_input(
              "Cantidad de etiquetas a imprimir",
              min_value=1,
              max_value=50,
              value=1,
          )

        zpl_code = ""
        for _ in range(int(cantidad_copias)):
          zpl_code += f"""
^XA
^FO30,25^ADN,30,15^FDASIAM AUTOPART^FS
^FO30,65^A0N,22,22^FDID: {id_val} | Asiam: {asiam_val}^FS
^FO30,95^A0N,20,20^FDMarca: {marca_val[:28]}^FS
^FO30,120^A0N,20,20^FDDESC: {prod_val[:30]}^FS
^FO30,145^A0N,20,20^FDOEM: {oem_val}^FS
^FO30,175^BY2,2,50^BCN,50,Y,N,N^FD{id_val}^FS
^XZ
"""

        st.text_area("Código ZPL listo para Zebra:", zpl_code, height=180)

        # Botón para impresión directa física conectada al puerto de Windows o red local
        if st.button("🖨️ IMPRIMIR DIRECTAMENTE EN ZEBRA"):
          try:
            # Envío directo por winspool si está en Windows (impresora térmica local)
            import platform

            if platform.system() == "Windows":
              import win32print

              printer_name = win32print.GetDefaultPrinter()
              hPrinter = win32print.OpenPrinter(printer_name)
              try:
                win32print.StartDocPrinter(
                    hPrinter, 1, ("Etiqueta ZPL ASIAM", None, "RAW")
                )
                win32print.StartPagePrinter(hPrinter)
                win32print.WritePrinter(hPrinter, zpl_code.encode("utf-8"))
                win32print.EndPagePrinter(hPrinter)
                win32print.EndDocPrinter(hPrinter)
                st.success(
                    f"✅ ¡Trabajo enviado con éxito a la impresora predeterminada:"
                    f" {printer_name}!"
                )
              finally:
                win32print.ClosePrinter(hPrinter)
            else:
              st.success(
                  "✅ Código generado. Conéctese desde el PC local de bodega con"
                  " la impresora configurada."
              )
          except Exception as print_err:
            st.warning(
                "⚠️ No se encontró una impresora predeterminada directa por"
                f" puerto Windows ({print_err}). Puede usar la descarga ZPL o"
                " verificar la conexión."
            )
            st.download_button(
                label="📥 Descargar archivo ZPL de respaldo",
                data=zpl_code,
                file_name=f"etiqueta_{id_val}.zpl",
                mime="text/plain",
            )
  else:
    st.warning("No hay datos cargados para generar etiquetas.")
